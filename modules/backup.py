"""Copia di sicurezza del database e fogli Excel leggibili da consegnare al commercialista."""
import datetime
import io

import pandas as pd

# Il modello Word delle distinte è salvato come testo base64 di alcuni MB: non sta in una cella Excel
# (massimo 32.767 caratteri) e si scarica a parte dalla pagina delle distinte.
TABELLE_ESCLUSE = {"modelli_documenti", "utenti", "accessi_ricordati"}  # utenti (password cifrate) e dispositivi ricordati non escono dall'app e il ripristino non li tocca
TABELLE_NON_RIPRISTINATE = {"registro_attivita"}  # nella copia per consultazione, ma il ripristino non cancella la storia
LIMITE_CELLA = 32767

# Fogli leggibili: nome del foglio -> query con due parametri (inizio e fine dell'esercizio)
FOGLI_ESERCIZIO = {
    "Prima nota": """
        SELECT m.data_registrazione AS Data, m.numero_documento AS Documento, m.causale AS Causale,
               m.tipo_attivita AS Attività, p.codice AS Conto, p.nome AS "Nome conto", r.descrizione AS Descrizione,
               r.dare AS "Dare (€)", r.avere AS "Avere (€)"
        FROM righe_prima_nota r
        JOIN movimenti_prima_nota m ON m.id = r.movimento_id
        JOIN piano_dei_conti p ON p.id = r.sottoconto_id
        WHERE m.data_registrazione BETWEEN ? AND ?
        ORDER BY m.data_registrazione, m.id, r.id
    """,
    "Totali per conto": """
        SELECT p.codice AS Conto, p.nome AS "Nome conto", p.tipo AS Tipo,
               SUM(r.dare) AS "Totale dare (€)", SUM(r.avere) AS "Totale avere (€)",
               SUM(r.dare) - SUM(r.avere) AS "Saldo dare-avere (€)"
        FROM righe_prima_nota r
        JOIN movimenti_prima_nota m ON m.id = r.movimento_id
        JOIN piano_dei_conti p ON p.id = r.sottoconto_id
        WHERE m.data_registrazione BETWEEN ? AND ?
        GROUP BY p.codice, p.nome, p.tipo
        ORDER BY p.codice
    """,
    "Ricevute": """
        SELECT r.numero_ricevuta AS Numero, r.data_emissione AS Data, a.denominazione AS Persona,
               a.codice_fiscale AS "Codice fiscale", r.causale AS Causale, r.importo AS "Importo (€)",
               r.modalita_pagamento AS Pagamento, r.marca_da_bollo AS "Marca da bollo (€)",
               CASE WHEN r.credito_pregresso = 1 THEN 'SÌ' ELSE '' END AS "Quota di esercizi precedenti"
        FROM ricevute_istituzionali r
        JOIN anagrafiche a ON a.id = r.anagrafica_id
        WHERE r.data_emissione BETWEEN ? AND ?
        ORDER BY r.data_emissione, r.id
    """,
    "Fatture sponsor": """
        SELECT f.numero_fattura AS Numero, f.data_fattura AS Data, a.denominazione AS Sponsor,
               a.partita_iva AS "Partita IVA", f.oggetto_contratto AS Oggetto, f.imponibile AS "Imponibile (€)",
               f.aliquota_iva AS "Aliquota IVA (%)", f.iva_totale AS "IVA (€)", f.totale_fattura AS "Totale (€)",
               f.iva_da_versare_50 AS "IVA da versare (€)", f.note AS Note
        FROM fatture_sponsor_398 f
        JOIN anagrafiche a ON a.id = f.sponsor_id
        WHERE f.data_fattura BETWEEN ? AND ?
        ORDER BY f.data_fattura, f.id
    """,
    "Incassi sponsor": """
        SELECT i.data_incasso AS Data, a.denominazione AS Sponsor, i.importo AS "Importo (€)",
               i.modalita_pagamento AS Pagamento, i.riferimento AS Riferimento
        FROM incassi_sponsor i
        JOIN anagrafiche a ON a.id = i.sponsor_id
        WHERE i.data_incasso BETWEEN ? AND ?
        ORDER BY i.data_incasso, i.id
    """,
    "Spese": """
        SELECT f.data_doc AS Data, f.numero_doc AS Documento, a.denominazione AS Fornitore,
               f.causale_spesa AS Causale, f.categoria_spesa AS Categoria, f.imponibile AS "Imponibile (€)",
               f.iva AS "IVA (€)", f.totale AS "Totale (€)"
        FROM fatture_passive f
        JOIN anagrafiche a ON a.id = f.fornitore_id
        WHERE f.data_doc BETWEEN ? AND ?
        ORDER BY f.data_doc, f.id
    """,
    "Compensi collaboratori": """
        SELECT c.data_erogazione AS Data, a.denominazione AS Collaboratore, a.codice_fiscale AS "Codice fiscale",
               s.mansione AS Mansione, s.tipo_contratto AS Contratto, c.causale AS Causale,
               c.importo_lordo AS "Lordo (€)", c.ritenuta_inps AS "Ritenuta INPS (€)",
               c.ritenuta_irpef AS "Ritenuta IRPEF (€)", c.importo_netto AS "Netto (€)"
        FROM compensi_sportivi c
        JOIN collaboratori_sportivi s ON s.id = c.collaboratore_id
        JOIN anagrafiche a ON a.id = s.anagrafica_id
        WHERE c.data_erogazione BETWEEN ? AND ?
        ORDER BY c.data_erogazione, c.id
    """,
    "Rimborsi trasferta": """
        SELECT t.data_partita AS Data, a.denominazione AS Persona, t.incontro_calcio AS Incontro,
               t.luogo_trasferta AS Luogo, t.km_percorsi AS Km, t.tariffa_aci AS "Tariffa ACI (€/km)",
               t.spese_pie_di_lista AS "Spese a piè di lista (€)", t.totale_rimborso AS "Totale (€)"
        FROM rimborsi_trasferta t
        JOIN anagrafiche a ON a.id = t.anagrafica_id
        WHERE t.data_partita BETWEEN ? AND ?
        ORDER BY t.data_partita, t.id
    """,
}


def _cella_sicura(valore):
    """Testi troppo lunghi vengono tagliati; quelli che iniziano con «=» non devono diventare formule Excel."""
    if isinstance(valore, str):
        if len(valore) > LIMITE_CELLA:
            valore = valore[:LIMITE_CELLA - 20] + " …[testo tagliato]"
        if valore.startswith("="):
            valore = "'" + valore
    return valore


def scrivi_foglio(writer, nome_foglio, df, riga_titolo=None):
    """Scrive il foglio con colonne già larghe quanto il contenuto e intestazione bloccata; il titolo facoltativo va nella prima riga."""
    df = df.map(_cella_sicura) if not df.empty else df
    riga_intestazione = 2 if riga_titolo else 0
    df.to_excel(writer, index=False, sheet_name=nome_foglio[:31], startrow=riga_intestazione)
    foglio = writer.sheets[nome_foglio[:31]]
    if riga_titolo:
        foglio.cell(row=1, column=1, value=_cella_sicura(riga_titolo)).font = foglio.cell(row=1, column=1).font.copy(bold=True, size=13)
    for i, colonna in enumerate(df.columns, start=1):
        lunghezze = [len(str(colonna))] + [len(str(v)) for v in df[colonna].head(200)]
        foglio.column_dimensions[foglio.cell(row=1, column=i).column_letter].width = min(max(lunghezze) + 2, 60)
    foglio.freeze_panes = f"A{riga_intestazione + 2}"


def elenco_tabelle(conn):
    righe = conn.execute(
        "SELECT name FROM sqlite_master WHERE type = 'table' AND name NOT LIKE 'sqlite_%' AND name NOT LIKE '_litestream%' ORDER BY name"
    ).fetchall()
    return [r[0] for r in righe if r[0] not in TABELLE_ESCLUSE]


def backup_completo(conn, versione_schema):
    """Excel con un foglio per ogni tabella del database, con tutte le righe e tutte le colonne."""
    buffer = io.BytesIO()
    tabelle = elenco_tabelle(conn)
    conteggi = []
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        dati = {t: pd.read_sql_query(f'SELECT * FROM "{t}"', conn) for t in tabelle}
        for t in tabelle:
            conteggi.append({"Tabella": t, "Righe": len(dati[t])})
        info = pd.DataFrame(
            [{"Tabella": "Creato il", "Righe": datetime.datetime.now().strftime("%d/%m/%Y %H:%M")},
             {"Tabella": "Versione dello schema", "Righe": versione_schema}] + conteggi)
        scrivi_foglio(writer, "Info", info)
        for t in tabelle:
            scrivi_foglio(writer, t, dati[t])
    return buffer.getvalue(), conteggi


def esportazione_esercizio(conn, inizio, fine, etichetta):
    """Excel leggibile dell'esercizio: prima nota, totali per conto, ricevute, sponsor, spese, compensi e rimborsi."""
    buffer = io.BytesIO()
    conteggi = []
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        fogli = {nome: pd.read_sql_query(sql, conn, params=[str(inizio), str(fine)]) for nome, sql in FOGLI_ESERCIZIO.items()}
        conteggi = [{"Foglio": nome, "Righe": len(df)} for nome, df in fogli.items()]
        info = pd.DataFrame([{"Foglio": "Esercizio", "Righe": etichetta},
                             {"Foglio": "Periodo", "Righe": f"{inizio.strftime('%d/%m/%Y')} - {fine.strftime('%d/%m/%Y')}"},
                             {"Foglio": "Creato il", "Righe": datetime.datetime.now().strftime("%d/%m/%Y %H:%M")}] + conteggi)
        scrivi_foglio(writer, "Info", info)
        for nome, df in fogli.items():
            scrivi_foglio(writer, nome, df)
    return buffer.getvalue(), conteggi


# ---------------------------------------------------------
# RIPRISTINO DA UNA COPIA COMPLETA
# ---------------------------------------------------------
def leggi_copia(file_bytes):
    """Legge un file creato da backup_completo(). Restituisce (versione dello schema, {tabella: (colonne, righe)}).
    Solleva ValueError se il file non è una copia completa dell'app."""
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as ex:
        raise ValueError(f"Il file non è un Excel leggibile ({ex}).")
    if "Info" not in wb.sheetnames:
        raise ValueError("Il file non è una copia completa creata dall'app: manca il foglio «Info».")
    versione = None
    for riga in wb["Info"].iter_rows(min_row=2, values_only=True):
        if riga and riga[0] == "Versione dello schema":
            try:
                versione = int(riga[1])
            except (TypeError, ValueError):
                pass
    if versione is None:
        raise ValueError("Nel foglio «Info» manca la versione dello schema: il file non sembra una copia completa dell'app.")
    tabelle = {}
    for nome in wb.sheetnames:
        if nome == "Info":
            continue
        righe = list(wb[nome].iter_rows(values_only=True))
        if not righe:
            continue
        colonne = [str(c) for c in righe[0] if c is not None]
        dati = []
        for r in righe[1:]:
            valori = list(r[:len(colonne)])
            if all(v is None for v in valori):
                continue
            # Testi che iniziano con «=» sono stati salvati con un apostrofo davanti per non diventare formule
            valori = [v[1:] if isinstance(v, str) and v.startswith("'=") else v for v in valori]
            dati.append(valori)
        tabelle[nome] = (colonne, dati)
    return versione, tabelle


def _sql_valore(v):
    if v is None or (isinstance(v, float) and v != v):
        return "NULL"
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, (int, float)):
        return repr(v)
    if isinstance(v, (datetime.datetime, datetime.date)):
        v = v.isoformat(sep=" ") if isinstance(v, datetime.datetime) else v.isoformat()
    return "'" + str(v).replace("'", "''") + "'"


def script_ripristino(conn, tabelle):
    """Script SQL che sostituisce il contenuto delle tabelle con quello della copia, in un'unica transazione.

    Si ripristinano solo le tabelle e le colonne presenti sia nella copia sia nel database; le tabelle escluse dalla copia
    (es. il modello Word delle distinte) restano come sono. Restituisce (script, {tabella: righe}).
    """
    esistenti = set(elenco_tabelle(conn)) - TABELLE_NON_RIPRISTINATE
    istruzioni = ["PRAGMA foreign_keys = OFF;", "BEGIN;"]
    if "esercizi_chiusi" in esistenti:
        istruzioni.append('DELETE FROM "esercizi_chiusi";')  # i blocchi tornano alla fine, con la prima nota già sostituita
    conteggi = {}
    for nome, (colonne, righe) in sorted(tabelle.items(), key=lambda t: t[0] == "esercizi_chiusi"):
        if nome not in esistenti:
            continue
        info = conn.execute(f'PRAGMA table_info("{nome}")').fetchall()
        colonne_db = {r[1]: r for r in info}  # (cid, nome, tipo, notnull, default, pk)
        indici = [i for i, c in enumerate(colonne) if c in colonne_db]
        nomi = [colonne[i] for i in indici]
        istruzioni.append(f'DELETE FROM "{nome}";')
        valori_righe = []
        for r in righe:
            valori = []
            for i in indici:
                v = r[i] if i < len(r) else None
                col = colonne_db[colonne[i]]
                if v is None and col[3]:
                    # Colonna obbligatoria lasciata vuota nel file (Excel non distingue testo vuoto da cella vuota)
                    v = "" if "TEXT" in str(col[2]).upper() or not col[2] else 0
                valori.append(_sql_valore(v))
            valori_righe.append("(" + ", ".join(valori) + ")")
        elenco_colonne = ", ".join(f'"{c}"' for c in nomi)
        for inizio in range(0, len(valori_righe), 200):
            istruzioni.append(f'INSERT INTO "{nome}" ({elenco_colonne}) VALUES ' + ", ".join(valori_righe[inizio:inizio + 200]) + ";")
        conteggi[nome] = len(righe)
    istruzioni += ["COMMIT;", "PRAGMA foreign_keys = ON;"]
    return "\n".join(istruzioni), conteggi
